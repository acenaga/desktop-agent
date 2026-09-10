"""
scripts/generate_uc01_fixtures.py: Genera documentos sintéticos para UC-01
con diferencias conocidas de fecha, alcance y monto (Sección 14.2).
"""

from pathlib import Path
import docx
from pypdf import PdfWriter


def create_fixtures():
    fixtures_dir = Path("fixtures/uc01")
    fixtures_dir.mkdir(parents=True, exist_ok=True)

    # 1. Propuesta A: Markdown
    doc_a = fixtures_dir / "propuesta_alfa.md"
    doc_a.write_text(
        """# Propuesta de Desarrollo de Software - Opción Alfa

**Fecha de entrega:** 15 de marzo de 2026  
**Empresa proveedora:** Soluciones Digitales Alfa S.A.  
**Monto total presupuestado:** $45.000 USD (cuarenta y cinco mil dólares)  

## Alcance del Proyecto
1. Desarrollo de interfaz de usuario para escritorio y web.
2. Desarrollo del backend con API REST para integración local.
3. Pruebas unitarias básicas de los módulos principales.
4. Periodo de garantía: 30 días posteriores a la entrega.

## Plazos
- Fase de diseño: 2 semanas.
- Fase de desarrollo: 6 semanas.
- Entrega final: 15 de marzo de 2026.
""",
        encoding="utf-8",
    )

    # 2. Propuesta B: DOCX usando python-docx
    doc_b_path = fixtures_dir / "propuesta_beta.docx"
    doc_b = docx.Document()
    doc_b.add_heading("Propuesta de Desarrollo de Software - Opción Beta", 0)
    
    p1 = doc_b.add_paragraph()
    p1.add_run("Fecha de entrega: ").bold = True
    p1.add_run("28 de abril de 2026\n")
    p1.add_run("Empresa proveedora: ").bold = True
    p1.add_run("Innovaciones Tecnológicas Beta SpA\n")
    p1.add_run("Monto total presupuestado: ").bold = True
    p1.add_run("$62.000 USD (sesenta y dos mil dólares)")

    doc_b.add_heading("Alcance del Proyecto", level=1)
    doc_b.add_paragraph(
        "1. Desarrollo integral de frontend (React + TypeScript) y backend (Python).\n"
        "2. Soporte para automatización nativa de escritorio y navegador dedicado.\n"
        "3. Despliegue en infraestructura de alta disponibilidad y CI/CD automatizado.\n"
        "4. Servicio de soporte técnico y mantenimiento 24/7 durante 12 meses."
    )

    doc_b.add_heading("Plazos y Cronograma", level=1)
    doc_b.add_paragraph(
        "Fase de arquitectura y prototipo: 3 semanas.\n"
        "Desarrollo e integraciones: 8 semanas.\n"
        "Certificación de seguridad y entrega final: 28 de abril de 2026."
    )
    doc_b.save(str(doc_b_path))

    print(f"[+] Fixtures generados en {fixtures_dir.resolve()}:")
    print(f"  - {doc_a.name} (Markdown)")
    print(f"  - {doc_b_path.name} (DOCX)")


if __name__ == "__main__":
    create_fixtures()
